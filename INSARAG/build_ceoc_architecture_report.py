from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "LinkGuard_INSARAG台灣CEOC開發架構查核報告_2026-05-22.docx"

FONT = "Microsoft JhengHei"
ACCENT = RGBColor(31, 78, 121)
DARK = RGBColor(11, 37, 69)
MUTED = RGBColor(85, 85, 85)

SOURCES = [
    (
        "中央災害應變中心作業要點修正資訊",
        "行政院中央災害防救會報",
        "https://cdprc.ey.gov.tw/Page/102B58FEB2E0F261/1c638843-4ebd-47cf-8269-6a79cbd905ad",
        "CEOC 作業要點修正、前置情資研判、進駐機關與數位科技任務。",
    ),
    (
        "中央災害應變中心作業要點 PDF",
        "行政院中央災害防救會報",
        "https://cdprc.ey.gov.tw/File/EE2774B9AD28C89A?A=C",
        "CEOC 開設、任務編組、災害類別與跨部會作業依據。",
    ),
    (
        "應變管理資訊系統 EMIC 介紹",
        "內政部消防署",
        "https://www.nfa.gov.tw/cht/index.php?act=article&article_id=10361&code=print&ids=1537",
        "災情查報、任務指派、通報傳送、指揮官決策支援。",
    ),
    (
        "國家災害防救科技中心 NCDR",
        "國家災害防救科技中心",
        "https://ncdr.nat.gov.tw/",
        "災害風險、情資整合、科技輔助決策與災防資料服務。",
    ),
    (
        "環境災害管理資訊系統",
        "環境部",
        "https://newemis.moenv.gov.tw/",
        "環境災害通報與中央災害應變中心相關資訊系統參考。",
    ),
    (
        "INSARAG Guidelines 2020",
        "INSARAG / UN OCHA",
        "https://insarag.org/methodology/insarag-guidelines/",
        "大型結構倒塌、國際 USAR、UCC、Information Management、Operational Field Guide。",
    ),
    (
        "INSARAG Information Management / ICMS",
        "INSARAG Information Management Working Group",
        "https://insarag.org/guidance-notes/manuals/information-management/",
        "ICMS RDC/UCC/Team Guide，作為 LinkGuard 災情與隊伍資料流參考。",
    ),
    (
        "INSARAG IEC/R",
        "INSARAG",
        "https://insarag.org/iec/iec-r/",
        "USAR team capability classification，支援隊伍能力表與派遣邏輯。",
    ),
    (
        "INSARAG USAR Coordination Manual",
        "INSARAG / UN OCHA",
        "https://insarag.org/wp-content/uploads/2017/06/UCC_Manual_v_22_Aug_-ilovepdf-compressed_1.pdf",
        "UCC/SCC/RDC/OSOCC coordination cycle、資訊收集分析與 briefing。",
    ),
    (
        "INSARAG Operations TRL",
        "INSARAG",
        "https://insarag.org/technical-reference-library/operations/",
        "USAR operations phase：RDC arrival、UCC/OSOCC registration、LEMA/NDMA coordination。",
    ),
    (
        "FEMA ICS Toolkit",
        "FEMA Preparedness Toolkit",
        "https://preptoolkit.fema.gov/web/nims-toolkit/ics",
        "ICS doctrine、forms、job aids、position-specific references。",
    ),
    (
        "FEMA ICS Forms",
        "FEMA Emergency Management Institute",
        "https://training.fema.gov/emiweb/is/icsresource/icsforms/",
        "ICS 201/202/203/204/205/206/207/208/209/213/214 等 IAP 與 incident management 表單。",
    ),
    (
        "WHO EMT minimum standards",
        "World Health Organization",
        "https://iris.who.int/handle/10665/341857",
        "Emergency Medical Team classification and minimum standards，支援醫療資料邊界。",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def width_to_inches(dxa: int):
    return Inches(dxa / 1440)


def set_table_widths(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = width_to_inches(widths[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_para(doc: Document, text: str = "", bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=DARK)
        rest = p.add_run(text[len(bold_prefix) :])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, color=ACCENT if level < 3 else DARK)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(item)
        set_run_font(run, 10.5)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(item)
        set_run_font(run, 10.5)


def add_callout(doc: Document, title: str, body: str, fill: str = "F4F7FB") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + " ")
    set_run_font(r, 10.5, True, DARK)
    r2 = p.add_run(body)
    set_run_font(r2, 10.5)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        set_cell_shading(cell, "EAF1F8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, 9.5, True, DARK)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            if c_idx == 0:
                set_cell_shading(cell, "F7F9FC")
            r = p.add_run(text)
            set_run_font(r, 9.2, c_idx == 0, DARK if c_idx == 0 else None)
            if c_idx in (0, len(row) - 1) and len(text) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in [("Heading 1", 15.5), ("Heading 2", 12.5), ("Heading 3", 11.2)]:
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT if style_name != "Heading 3" else DARK
        style.paragraph_format.space_before = Pt(11 if style_name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("LinkGuard INSARAG × 台灣 CEOC 開發架構查核報告｜2026-05-22")
    set_run_font(r, 8.5, color=MUTED)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("LinkGuard INSARAG × 台灣 CEOC")
    set_run_font(r, 22, True, DARK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    r2 = p2.add_run("開發架構查核報告")
    set_run_font(r2, 20, True, ACCENT)
    p3 = doc.add_paragraph()
    r3 = p3.add_run("依 INSARAG / FEMA ICS-NIMS / WHO EMT / 台灣 CEOC-EMIC 官方資料整理｜2026-05-22")
    set_run_font(r3, 10.5, color=MUTED)

    add_callout(
        doc,
        "總結判斷：",
        "LinkGuard v0.3 的 role-based apps、shared core、offline-first、AAR 與 medical separation 方向正確；若要落地台灣，下一步必須把 CEOC/EMIC-style 災情通報、任務派遣、跨機關通報、地方應變中心與現場指揮所的資料邊界補成正式產品契約。",
    )


def section_summary(doc: Document) -> None:
    add_heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "本報告把 LinkGuard v0.3 放進兩個真實架構中檢查：第一是 INSARAG 國際 USAR 方法論，第二是台灣 CEOC 中央災害應變中心與 EMIC 類災情任務體系。結論不是要把 LinkGuard 變成政府既有系統的複製品，而是把資料模型、角色權限、離線同步與現場回報做成能與真實應變體系對話的產品。",
    )
    add_bullets(
        doc,
        [
            "已正確：UCC/SCC/TL/TE/VO/EMT 分流、shared incident state、離線佇列、AAR、醫療資料邊界。",
            "要補強：CEOC/EMIC-compatible 災情與任務資料模型、正式同步服務、裝置 provisioning、ASR/RCM/worksite 產品化。",
            "不能宣稱：目前不能說已完整接軌台灣 CEOC/EMIC，也不能說是完整 INSARAG field-ready 系統。",
            "開發優先序：資料契約先於 UI 美化；同步與稽核先於 AI；醫療摘要邊界先於跨角色資料開放。",
        ],
    )
    add_callout(
        doc,
        "核心架構原則：",
        "CEOC/UCC 看跨區與跨機關協調，SCC/地方應變中心掌握現場戰術，TL/TE/VO 降低回報負擔，EMT 保有臨床資料權威。所有角色共享同一個 incident state，但不能共享同一個資料權限。",
        "EEF6EF",
    )


def section_sources(doc: Document) -> None:
    add_heading(doc, "2. 官方基準與查核邏輯", 1)
    add_para(
        doc,
        "本報告使用官方一手或高可信官方資料作為查核基準。國際部分以 INSARAG、FEMA ICS/NIMS、WHO EMT 為主；台灣部分以行政院中央災害防救會報、內政部消防署 EMIC、NCDR 與環境部系統為主。",
    )
    add_table(
        doc,
        ["來源", "本報告採用重點", "對 LinkGuard 的要求"],
        [
            ["台灣 CEOC 作業要點", "開設、前置情資研判、進駐機關、任務分工", "UCC 需支援跨機關、跨行政區、開設層級與會議紀錄。"],
            ["消防署 EMIC", "災情查報、任務指派、通報傳送、決策支援", "建立 DisasterReport、Mission、AgencyMessage 與查證狀態。"],
            ["INSARAG Guidelines / ICMS", "UCC/SCC/RDC、資訊管理、USAR coordination", "建立 team intake、worksite、ASR、RCM 與 coordination cycle。"],
            ["FEMA ICS / IAP forms", "目標、組織、任務、通訊、醫療、安全、活動紀錄", "建立 IAP-style operational period 與 audit trail。"],
            ["WHO EMT", "EMT 分類、最低標準、醫療協調", "醫療資料分 summary 與 clinical details，不可全域暴露。"],
        ],
        [1950, 3600, 3810],
    )


def section_ceoc(doc: Document) -> None:
    add_heading(doc, "3. 台灣 CEOC 體系與產品定位", 1)
    add_para(
        doc,
        "CEOC 是台灣中央災害應變中心的核心縮寫。本報告使用 CEOC 指稱台灣中央層級的災害應變協調中心；EOC 僅作為一般 Emergency Operations Center 或地方應變中心概念，不取代 CEOC。",
    )
    add_para(
        doc,
        "LinkGuard-UCC 不應直接宣稱等同 CEOC。較準確的產品定位是：LinkGuard-UCC 可提供 CEOC-like 的跨區協調、資料彙整、跨機關通報、情資視覺化與決策支援能力；LinkGuard-SCC 則對應地方災害應變中心、前進指揮所或特定災區的現場戰術節點。",
    )
    add_table(
        doc,
        ["台灣節點", "LinkGuard 對應", "主要資料", "產品邊界"],
        [
            ["CEOC", "UCC", "跨縣市災情、資源、會議、跨機關通報", "看全局與協調，不直接覆蓋現場戰術資料。"],
            ["地方災害應變中心", "SCC", "行政區災情、派遣、現場回報、地方資源", "管理地方/災區戰術與任務落地。"],
            ["前進指揮所", "SCC / TL", "現場分區、隊伍、worksite、危險與通訊", "負責可執行任務與現場安全。"],
            ["查報人員", "TE", "GPS、照片、災情狀態、任務回報", "低負擔回報，不進入完整指揮後台。"],
            ["民眾/志工", "VO", "災情通報、位置、照片、SOS、安全需求", "需查證，不直接成為官方 operational truth。"],
            ["醫療組/EMT", "EMT", "檢傷、後送、容量、交接摘要", "臨床細節保留在醫療權限邊界。"],
        ],
        [1450, 1300, 3300, 3310],
    )


def section_role_architecture(doc: Document) -> None:
    add_heading(doc, "4. LinkGuard 角色架構對照", 1)
    add_para(
        doc,
        "角色分流是 LinkGuard v0.3 最重要的正確方向。真實災害應變不是所有人看同一個 HQ 畫面，而是每個角色只看到自己能決策、能執行、能回報的資訊。",
    )
    add_table(
        doc,
        ["App/角色", "台灣 CEOC 場景", "INSARAG/ICS 場景", "應開放能力"],
        [
            ["UCC", "CEOC-like 協調層", "UCC/OSOCC coordination + Unified Command", "跨區 dashboard、資源缺口、會議、AAR、資料權威。"],
            ["SCC", "地方應變中心/現場指揮所", "SCC/ICP", "分區、任務派遣、人員安全、地圖態勢、現場同步。"],
            ["TL", "前進指揮/分隊長", "Division/Group/Team Leader", "隊伍任務、worksite 更新、危險標記、進出場。"],
            ["TE", "查報/搜救隊員", "Field operator", "GPS、照片、任務狀態、ASR/RCM 線索、SOS。"],
            ["VO", "民眾/志工通報端", "Volunteer/community reporter", "低門檻通報、語音/照片/位置、查證前狀態。"],
            ["EMT", "醫療組/救護單位", "Emergency Medical Team", "START、生命徵象、後送、醫療容量、交接摘要。"],
        ],
        [1050, 2200, 2550, 3560],
    )
    add_callout(
        doc,
        "不可越權：",
        "UCC 不能直接改寫 SCC 的現場戰術 truth；VO/TE 回報不能直接變成官方災情；SCC/UCC 不應讀取 EMT 完整臨床病歷。這三條是產品資料權限的底線。",
        "FFF4E5",
    )


def section_development_blueprint(doc: Document) -> None:
    add_heading(doc, "5. 開發架構藍圖", 1)
    add_para(
        doc,
        "建議以共同 incident state 為核心，外層切成 role-based view。每個寫入事件都需要 actor、role、device、timestamp、source、sync state 與 audit event。這讓 LinkGuard 同時支援現場離線、跨中心同步、AAR 與事後稽核。",
    )
    add_table(
        doc,
        ["架構層", "責任", "最低資料契約"],
        [
            ["Shared Incident State", "所有角色共同指向的事件真相來源", "incidentID、災害類別、行政區、operational period、版本、資料權威。"],
            ["Role View", "依 UCC/SCC/TL/TE/VO/EMT 裁切 UI 與功能", "role、permission、read scope、write scope、approval path。"],
            ["Domain Authority", "每個資料域有明確維護者", "command、field operation、medical、resource、communications、AAR。"],
            ["Offline Queue", "斷線時保存操作，恢復後同步", "operationID、payload、retry、receipt、merge status。"],
            ["Audit Trail", "支援 AAR、責任追蹤與資料稽核", "actor、time、location、before/after、source command。"],
            ["CEOC/EMIC Adapter", "把產品資料映射成台灣災情與任務語意", "DisasterReport、Mission、AgencyMessage、ResourceRequest。"],
        ],
        [1900, 3250, 4210],
    )
    add_heading(doc, "5.1 安裝與模組啟用架構", 2)
    add_para(
        doc,
        "LinkGuard 不應把 UCC、SCC、TL、TE、VO、EMT 做成彼此完全分裂的多個產品，也不應把所有功能塞進同一個混亂首頁。較合理的產品架構是：行動裝置採單一 App，登入後依帳號權限啟用角色模組；指揮端採 Command Launcher，登入後依組織、事件與職務開啟 CEOC/UCC/SCC/地圖/任務/醫療/AAR 等工作模組。",
    )
    add_callout(
        doc,
        "產品定義：",
        "安裝只是部署入口，真正決定使用者看到什麼的是帳號權限、組織歸屬、事件範圍、任務角色與模組授權。這可避免現場裝錯版本，也讓所有人能使用同一套版本更新。",
        "EEF6EF",
    )
    add_table(
        doc,
        ["產品型態", "適用裝置", "啟用方式", "主要模組"],
        [
            ["LinkGuard Mobile", "iPhone / iPad / Android field device", "登入帳號、QR code、邀請碼或 MDM provisioning 後啟用", "VO 通報、TE 任務、TL 分隊、EMT 醫療、SCC Mobile。"],
            ["LinkGuard Command", "Mac / iPad command workstation", "帳號 + 組織 + incident scope + command authority", "CEOC/UCC Dashboard、SCC Command、Map Operations、AgencyMessage、Resource、AAR。"],
            ["Admin / Provisioning", "管理端或受控工作站", "系統管理員、機關管理員或事件管理員授權", "帳號、角色、裝置、事件、模組 entitlement、撤銷與臨時授權。"],
        ],
        [1750, 2350, 3050, 2210],
    )
    add_table(
        doc,
        ["登入後權限", "會啟用的模組", "不應顯示"],
        [
            ["CEOC/UCC", "跨區災情、AgencyMessage、資源缺口、會議、AAR、CEOC dashboard", "TE 詳細任務操作、完整臨床病歷。"],
            ["SCC", "地方/現場災情查證、任務派遣、地圖態勢、同步節點、人員安全", "跨區戰略覆寫、完整臨床病歷。"],
            ["TL", "分區任務、worksite、ASR/RCM、隊伍狀態、安全進出", "CEOC 會議、跨機關資源決策。"],
            ["TE", "任務接收、GPS、照片、ASR/RCM 回報、SOS、簡短通訊", "指揮後台、其他隊伍敏感位置。"],
            ["VO", "低門檻災情通報、位置、照片、狀態查詢、安全提示", "ICS/CEOC 指揮、任務派遣、個資資料。"],
            ["EMT", "傷患、START、生命徵象、後送、醫療容量、交接摘要", "非醫療指揮後台與不必要現場個資。"],
        ],
        [1500, 5100, 2760],
    )
    add_bullets(
        doc,
        [
            "安裝包統一：同一個版本更新下去，使用者登入後自動看到自己的模組。",
            "權限動態：救災現場角色會變，server/provisioning service 必須可臨時升級、降級、撤銷或轉移權限。",
            "離線保護：斷網時使用最後有效權限快照，但高風險操作需排隊，恢復連線後再取得確認或合併。",
            "訓練簡化：現場只需知道「安裝 LinkGuard、登入帳號、進入自己的首頁」，不要要求使用者判斷該裝哪個版本。",
            "部署彈性：CEOC/UCC/SCC 可用 Command Launcher，行動端維持 field-first 低負擔操作。",
        ],
    )
    add_heading(doc, "5.2 CEOC-compatible data model", 2)
    add_table(
        doc,
        ["模型", "必要欄位", "權威角色"],
        [
            ["Incident", "災害類別、開設層級、主管機關、行政區、作戰週期", "UCC/SCC"],
            ["DisasterReport", "來源、類型、位置、照片、狀態、查證層級、是否派工", "SCC；UCC 彙整"],
            ["Mission", "派遣單位、任務類型、優先度、接收/執行/完成/退回", "SCC/TL"],
            ["AgencyMessage", "跨機關通報、工作會報、重要情資、回覆確認", "UCC/SCC"],
            ["ResourceRequest", "人力、車輛、機具、醫療、物資、避難收容需求", "UCC/SCC/Logistics"],
            ["MedicalSummary", "傷患數、檢傷分類、後送狀態、醫療容量摘要", "EMT"],
            ["AuditEvent", "誰、何時、在哪裡、依哪個權限改了什麼", "System"],
        ],
        [1750, 5250, 2360],
    )
    add_heading(doc, "5.3 sync envelope and approval path", 2)
    add_para(
        doc,
        "CEOC/INSARAG 對齊不是把資料送上雲端就完成。每筆資料都需要知道來源、權限、查證狀態與是否已被納入 official operational picture。這一層是 LinkGuard 從展示 demo 進入真實救災系統的關鍵。",
    )
    add_table(
        doc,
        ["欄位", "用途", "最低行為"],
        [
            ["operationID", "每次操作的穩定識別碼", "重送時不可重複建立災情或任務。"],
            ["sourceRole/sourceDevice", "標記由誰、哪個裝置產生", "支援離線稽核與責任追蹤。"],
            ["domain", "command、field、medical、resource、agency、AAR", "依 domain 決定可讀寫角色。"],
            ["approvalState", "draft、submitted、accepted、rejected、superseded", "VO/TE 回報必須經查證才成為 official truth。"],
            ["syncReceipt", "queued、sent、acknowledged、merged、conflicted", "UCC/SCC 能看到資料是否真的同步完成。"],
            ["auditPointer", "連到 AAR/audit event", "每個狀態轉換都能回放。"],
        ],
        [1750, 3350, 4260],
    )
    add_callout(
        doc,
        "實作建議：",
        "先建立 envelope 與 approvalState，再做漂亮 dashboard。若沒有這層，CEOC/EMIC 災情、INSARAG ASR/RCM、醫療摘要會彼此覆蓋，最後無法判斷哪一筆是正式狀態。",
        "EEF6EF",
    )


def section_taiwan_data_flow(doc: Document) -> None:
    add_heading(doc, "6. 台灣災情資料流", 1)
    add_para(
        doc,
        "EMIC 類系統的核心不是地圖本身，而是災情、任務、查報與通報的閉環。LinkGuard 要落地台灣，必須讓 VO/TE/TL/SCC/UCC 的資料狀態能對應到通報、查證、派工、執行、完成、彙整與會議決策。",
    )
    add_numbered(
        doc,
        [
            "民眾/志工或現場人員建立 DisasterReport：包含位置、照片、類型、描述與安全狀態。",
            "SCC/地方應變中心進行查證：標記為待查證、已查證、重複、退回、已派工。",
            "SCC 建立 Mission：指派單位、任務類型、優先度、回報要求與完成條件。",
            "TL/TE 執行任務並回報：狀態、GPS、照片、語音、危險、需求與完成證據。",
            "UCC/CEOC-like 視圖彙整跨行政區資訊：災情趨勢、資源缺口、重大事件、跨機關通報。",
            "AgencyMessage 形成工作會報與重要情資紀錄：保留發送、回覆確認與責任單位。",
            "AuditEvent 自動記錄所有狀態變更，供 AAR 與事後稽核。",
        ],
    )
    add_table(
        doc,
        ["狀態", "產品意義", "不可省略原因"],
        [
            ["Submitted", "前線或民眾已送出", "保留來源，不讓資料消失。"],
            ["Verifying", "SCC/權責單位查證中", "避免未查證資料直接變成官方災情。"],
            ["Accepted", "納入 operational truth", "地圖、任務、會議與統計可引用。"],
            ["Dispatched", "已形成任務", "把災情轉為可執行救災行動。"],
            ["In Progress", "任務進行中", "指揮端可掌握現場狀態。"],
            ["Completed", "任務完成並附證據", "支援結案、統計與 AAR。"],
            ["Rejected/Duplicate", "退回或重複", "保留紀錄但不污染官方態勢。"],
        ],
        [1500, 3650, 4210],
    )
    add_heading(doc, "6.1 CEOC/EMIC-style workflow owners", 2)
    add_table(
        doc,
        ["流程節點", "主要 owner", "LinkGuard UI/服務", "資料輸出"],
        [
            ["前置情資研判", "UCC / CEOC-like Planning", "情資 dashboard、預警、風險摘要", "incident draft、hot zone、resource pre-plan"],
            ["災情通報", "VO / TE / 外部來源", "低門檻通報、照片、定位、語音", "DisasterReport submitted"],
            ["查證與分派", "SCC / 地方應變中心", "查證佇列、地圖、派工板", "accepted report、Mission dispatched"],
            ["任務執行", "TL / TE / EMT", "任務卡、進度、照片、傷患摘要", "Mission update、MedicalSummary"],
            ["跨機關通報", "UCC / SCC", "AgencyMessage、工作會議、回覆確認", "會議紀錄、重要情資、責任單位"],
            ["結案與復盤", "UCC / SCC / AAR", "timeline replay、稽核、指標", "AAR package、改善項、訓練材料"],
        ],
        [1750, 2050, 3000, 2560],
    )


def section_insarag_flow(doc: Document) -> None:
    add_heading(doc, "7. INSARAG 作業資料流", 1)
    add_para(
        doc,
        "INSARAG 場景與台灣 CEOC 場景可以共用同一套資料骨架，但名稱與工作節奏不同。INSARAG 重視 USAR team capability、RDC arrival、UCC/SCC coordination、worksite、ASR、RCM 與 demobilization；台灣場景重視災情查報、任務派遣、跨機關通報與行政區彙整。",
    )
    add_table(
        doc,
        ["INSARAG object", "LinkGuard object", "開發要求"],
        [
            ["RDC arrival", "TeamIntake / ResourceCheckIn", "登錄隊伍、國家/單位、能力、抵達、支援需求。"],
            ["UCC/SCC", "UCC/SCC command scope", "支援 coordination cycle：collection、analysis、briefing、distribution。"],
            ["Sector", "Sector / SubSector", "地圖分區、責任單位、危險、active worksite。"],
            ["Worksite", "Worksite", "code、位置、ASR level、status、assigned team、照片。"],
            ["ASR", "Assessment", "危險、入口、受困者線索、優先度、reporter、timestamp。"],
            ["RCM", "Rescue/Marking feature", "現場標記、搜索狀態、victim marking、cordon、照片/GPS。"],
            ["Medical evacuation", "MedicalSummary / PatientFlow", "檢傷、CCP、後送路線、目的地、交接狀態。"],
            ["Demobilization", "Demobilization record", "撤收、移交、隊伍狀態、設備、人員與 AAR input。"],
        ],
        [1700, 2550, 5110],
    )
    add_callout(
        doc,
        "資料整合重點：",
        "同一個 worksite 必須能追到 sector、assigned team、ASR、RCM、task、medical event 與 AAR。若這條鏈斷掉，系統就只是多個漂亮頁面，不是真正的救災作業系統。",
        "F4F7FB",
    )
    add_heading(doc, "7.1 ASR / RCM / Worksite implementation minimum", 2)
    add_table(
        doc,
        ["功能", "最低欄位", "UI 行為"],
        [
            ["Worksite", "code、sector、location、priority、status、assigned team", "SCC 建立與指派；TL/TE 僅能更新被授權 worksite。"],
            ["ASR", "level、hazards、victim clues、access、reporter、evidence", "TE/TL 可離線提交；SCC 接受後更新官方狀態。"],
            ["RCM", "marking type、coordinates、photo、timestamp、team、status", "地圖可視化；保留歷史與 current operational marking。"],
            ["Team capability", "team code、classification、technical/dog/medical/hazmat、support needs", "UCC/SCC 用於派遣，不讓隊伍自動取得 command authority。"],
            ["Medical link", "patient count、START summary、CCP、evacuation route、handoff", "只回掛營運摘要，不把 clinical notes 暴露給全角色。"],
        ],
        [1700, 3900, 3760],
    )


def section_gap_matrix(doc: Document) -> None:
    add_heading(doc, "8. Gap Matrix", 1)
    add_para(
        doc,
        "以下矩陣把目前 LinkGuard v0.3 的可取方向、台灣 CEOC 缺口、INSARAG 缺口與開發優先級放在同一張表。這張表可直接作為後續 issue/roadmap 的來源。",
    )
    add_table(
        doc,
        ["面向", "目前可取方向", "缺口", "優先級"],
        [
            ["角色分流", "UCC/SCC/TL/TE/VO/EMT 已有清楚產品方向", "CEOC、地方應變中心、前進指揮所命名與權限仍需正式化", "P0"],
            ["資料模型", "shared core 與 incident state 已有雛形", "缺 CEOC/EMIC-style DisasterReport、Mission、AgencyMessage", "P0"],
            ["同步", "offline queue 與 retry 思路正確", "缺正式 server endpoint、device provisioning、receipt dashboard", "P0"],
            ["地圖", "點線面與 overlay 概念正確", "ASR/RCM/worksite 尚未完整產品化", "P1"],
            ["醫療", "EMT 獨立 app 與 medical separation 方向正確", "缺 operational summary vs clinical details 的強制資料邊界", "P1"],
            ["通報/會議", "已有 AAR 與 sync audit 思路", "缺工作會報、跨機關通報、回覆確認、重要情資 workflow", "P1"],
            ["AI", "可做風險與摘要輔助", "不能在資料契約未穩時先做決策自動化", "P2"],
            ["正式宣稱", "prototype/core scaffold 可成立", "尚未能宣稱 CEOC/EMIC integrated 或 INSARAG field-ready", "P0"],
        ],
        [1350, 2850, 3910, 1250],
    )
    add_heading(doc, "8.1 Development roadmap", 2)
    add_table(
        doc,
        ["階段", "目標", "完成後能力"],
        [
            ["P0 資料契約", "CEOC/EMIC models、INSARAG worksite/ASR/RCM、sync envelope", "所有角色可寫入同一 incident state，且每筆資料有權限與稽核。"],
            ["P0 同步底座", "server endpoint、device provisioning、receipt/conflict dashboard", "弱網與跨裝置不再只是 local prototype。"],
            ["P1 指揮產品", "UCC/SCC dashboard、查證佇列、任務板、AgencyMessage", "支援 CEOC-like 跨機關與 SCC 現場戰術作業。"],
            ["P1 現場產品", "TL/TE/VO/EMT 任務、照片、GPS、ASR/RCM、醫療摘要", "可進行現場演練，不靠手動補資料。"],
            ["P2 演練與 AAR", "field drill、timeline replay、指標、改善清單", "能用真實演練結果修正產品。"],
            ["P3 AI 輔助", "風險摘要、資源建議、重複災情合併", "在資料契約穩定後做輔助，不取代指揮官決策。"],
        ],
        [1450, 3250, 4660],
    )


def section_acceptance(doc: Document) -> None:
    add_heading(doc, "9. 開發驗收清單", 1)
    add_para(
        doc,
        "以下條件達成後，LinkGuard 才能比較合理地宣稱具備台灣 CEOC 對齊與 INSARAG 架構對齊的開發基礎。它不是完整正式認證，但可以作為產品開發 gate。",
    )
    add_bullets(
        doc,
        [
            "CEOC：UCC 可看跨行政區災情、資源、任務、工作會報與跨機關通報摘要。",
            "地方/SCC：SCC 可建立災情、查證、派工、追蹤、結案，且與 UCC 同步。",
            "現場/TL/TE：任務接收、GPS、照片、危險、ASR/RCM、SOS 可離線保存並恢復同步。",
            "VO：民眾/志工回報進入查證流程，不直接覆蓋官方態勢。",
            "EMT：傷患資料分成 clinical details 與 operational summary，SCC/UCC 只看必要摘要。",
            "INSARAG：RDC/team intake、UCC/SCC、sector/worksite、ASR、RCM、demobilization 形成閉環。",
            "同步：每筆操作有 receipt、merge status、conflict handling 與 audit event。",
            "AAR：可從時間線回放災情、任務、醫療、通報、同步與決策事件。",
        ],
    )
    add_table(
        doc,
        ["開發 gate", "通過標準", "失敗時不可宣稱"],
        [
            ["CEOC-compatible", "災情/任務/通報/資源/會議資料可被 UCC 彙整", "不可宣稱台灣 CEOC 對齊"],
            ["INSARAG-aligned", "worksite-ASR-RCM-team-task-medical-AAR 可追溯", "不可宣稱 INSARAG field-ready"],
            ["Offline-ready", "無網路完成回報與任務，恢復後可同步與稽核", "不可宣稱現場弱網可用"],
            ["Medical-safe", "臨床資料與營運摘要強制分離", "不可宣稱醫療資料權限安全"],
            ["Audit-ready", "每筆決策與資料改動可追到人、時間、來源", "不可宣稱可做正式 AAR"],
        ],
        [1850, 4700, 2810],
    )


def section_appendices(doc: Document) -> None:
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "Appendix A. 官方表單與 LinkGuard 模型對照", 1)
    add_para(
        doc,
        "本附錄把官方作業常見表單或資訊流轉成開發對照。目的不是複製每一張表，而是確保 LinkGuard 的資料模型可以承接真實應變需要。",
    )
    add_table(
        doc,
        ["官方表單/資料流", "LinkGuard 對應", "開發備註"],
        [
            ["CEOC 作業要點/工作會議", "AgencyMessage、MeetingRecord、DecisionLog", "支援跨機關通報、回覆確認、重要情資與決策紀錄。"],
            ["EMIC 災情/任務", "DisasterReport、Mission、ReportVerification", "災情不能直接等於任務，必須有查證與派工狀態。"],
            ["ICS 201 Incident Briefing", "IncidentBrief、InitialSituation", "事件初始摘要與轉移指揮時使用。"],
            ["ICS 202 Incident Objectives", "OperationalPeriod、Objectives", "每個 operational period 應有明確目標。"],
            ["ICS 203/207 Organization", "CommandStructure、RoleAssignment", "UCC/SCC/TL/EMT 權責需要可視化與可稽核。"],
            ["ICS 204 Assignment List", "Mission、TaskAssignment", "任務必須有 owner、目標、資源、回報要求。"],
            ["ICS 205/205A Communications", "CommunicationPlan、AgencyContact", "支援 PTT、群組、備援、聯絡清單。"],
            ["ICS 206 Medical Plan", "MedicalSummary、EvacuationPlan", "只同步營運摘要，臨床細節留在 EMT。"],
            ["ICS 214 Activity Log", "AuditEvent、AARTimeline", "所有角色行為轉為 AAR 可回放事件。"],
            ["INSARAG ICMS UCC/RDC/Team", "TeamIntake、UCCDashboard、RDCRegistration", "支援隊伍到達、能力、分派、協調與撤收。"],
        ],
        [2500, 2850, 4010],
    )
    add_heading(doc, "Appendix B. 演練情境驗收", 1)
    add_table(
        doc,
        ["情境", "操作路徑", "必須觀察到的結果"],
        [
            ["地震多點災情", "VO 通報三筆災情，SCC 查證兩筆並派工，UCC 看跨區摘要", "未查證資料不污染 UCC 統計；派工任務可追到原始災情。"],
            ["重型 USAR 抵達", "UCC 建立 TeamIntake，SCC 指派 worksite，TL/TE 回報 ASR/RCM", "team capability 影響派遣；ASR/RCM 可回掛 worksite。"],
            ["弱網現場任務", "TE 離線拍照與回報，恢復網路後同步", "operationID 不重複；sync receipt 可見；AAR 留下時間線。"],
            ["大量民眾通報", "VO 建立多筆照片/定位通報，SCC 合併重複", "duplicate/rejected 保留原始紀錄但不進 official truth。"],
            ["醫療後送", "EMT 建立 START 與後送摘要，SCC/UCC 查看容量與路線", "SCC/UCC 只看 MedicalSummary，不看完整 clinical details。"],
            ["跨機關會議", "UCC 發送 AgencyMessage，SCC/醫療/後勤回覆確認", "會議決策、責任單位、回覆狀態進入 audit trail。"],
        ],
        [1700, 3650, 4010],
    )
    add_heading(doc, "Appendix C. 不應做的設計", 1)
    add_bullets(
        doc,
        [
            "不要把 UCC 做成所有資料都能直接改的 super admin；這會破壞 SCC 現場戰術權。",
            "不要讓 VO/民眾通報直接變成官方災情；必須有查證與接受狀態。",
            "不要讓完整醫療病歷流到 UCC/SCC dashboard；只同步營運必要摘要。",
            "不要先做 AI 指揮建議再補資料契約；AI 應建立在可稽核資料流之上。",
            "不要把 INSARAG ASR/RCM 做成單純文字欄位；它們必須能連到 worksite、地圖、任務與 AAR。",
        ],
    )


def section_sources_list(doc: Document) -> None:
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "Sources", 1)
    add_para(doc, "以下為本報告使用的官方來源與用途。")
    for name, publisher, url, use in SOURCES:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{name}｜{publisher}：")
        set_run_font(r, 9.7, True, DARK)
        add_hyperlink(p, url, url)
        p.add_run(f"。用途：{use}")
        for run in p.runs:
            if run.text.startswith("。用途"):
                set_run_font(run, 9.5)


def build() -> Path:
    doc = Document()
    configure_doc(doc)
    add_cover(doc)
    section_summary(doc)
    section_sources(doc)
    section_ceoc(doc)
    section_role_architecture(doc)
    section_development_blueprint(doc)
    section_taiwan_data_flow(doc)
    section_insarag_flow(doc)
    section_gap_matrix(doc)
    section_acceptance(doc)
    section_appendices(doc)
    section_sources_list(doc)
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
