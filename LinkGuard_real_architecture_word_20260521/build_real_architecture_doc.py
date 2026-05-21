from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "LinkGuard_真實救災架構查核報告_2026-05-21.docx"


SOURCES = [
    {
        "name": "INSARAG Guidelines 2020",
        "publisher": "INSARAG / UN OCHA",
        "url": "https://insarag.org/methodology/insarag-guidelines/",
        "use": "USAR/INSARAG methodology, UCC, information management, operations field guide.",
    },
    {
        "name": "National Incident Management System, 3rd Edition",
        "publisher": "FEMA",
        "url": "https://www.fema.gov/sites/default/files/2020-07/fema_nims_doctrine-2017.pdf",
        "use": "ICS common hierarchy, Command/Operations/Planning/Logistics/Finance/Admin functions.",
    },
    {
        "name": "FEMA ICS Toolkit",
        "publisher": "FEMA Preparedness Toolkit",
        "url": "https://preptoolkit.fema.gov/web/nims-toolkit/ics",
        "use": "ICS doctrine, forms, guides, position job aids, and qualification resources.",
    },
    {
        "name": "FEMA Incident Action Planning",
        "publisher": "FEMA EMI",
        "url": "https://emilms.fema.gov/is_822/groups/139.html",
        "use": "IAP process, operational-period planning rhythm, resource assignments, ongoing situational awareness.",
    },
    {
        "name": "Classification and minimum standards for emergency medical teams",
        "publisher": "World Health Organization",
        "url": "https://www.who.int/publications/i/item/9789240029330/",
        "use": "EMT coordination with emergency response systems and minimum medical team standards.",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = width_to_inches(widths[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def width_to_inches(dxa: int):
    return Inches(dxa / 1440)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.style = doc.styles["Normal"]
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)


def add_numbered(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(4)


def add_callout(doc: Document, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    r = p.add_run(title + " ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 58, 95)
    p.add_run(body)
    doc.add_paragraph()


def add_matrix_table(doc: Document):
    rows = [
        ["版本", "真實對應", "主要責任", "不應承擔"],
        ["UCC", "跨區/國際協調與資料權威層", "多災區資源總覽、對外協調、跨中心同步、資訊管理、AAR 彙整", "不應直接取代現場戰術指揮"],
        ["SCC", "現場指揮中心/ICP 戰術層", "分區、任務、現場安全、人員狀態、地圖態勢、現場同步節點", "不應變成所有戰略、醫療、財務資料的唯一入口"],
        ["TL", "Division/Group/Team Leader", "接收任務、分派小隊、更新 worksite、回報進度與危險標記", "不應發布跨區戰略命令"],
        ["TE", "搜救隊員/任務執行者", "任務狀態、GPS、照片、危險、SOS、簡短回報", "不應看到完整醫療病歷或指揮後台"],
        ["VO", "志工/民眾支援回報端", "低複雜度災情回報、照片、位置、SOS、多語/語音輔助", "不應進入 ICS 指揮、分區建立、完整人員追蹤"],
        ["EMT", "Emergency Medical Team / Medical Branch", "檢傷、傷患卡、生命徵象、後送、交接、醫院容量", "不應被搜救任務管理與非醫療指揮流程干擾"],
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    widths = [1150, 2050, 3550, 2610]
    set_table_widths(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            if r_idx == 0:
                run.bold = True
                set_cell_shading(cell, "F2F4F7")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 0 and r_idx > 0:
                run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def add_data_table(doc: Document):
    rows = [
        ["資料域", "真實系統需要保存", "權威/主要維護者"],
        ["Incident / Operational Period", "事故 ID、地點、目標、作戰週期、IAP 版本、狀態摘要", "SCC 建立現場版；UCC 彙整跨區版"],
        ["Worksite / Sector Map", "分區、子分區、案場、危險區、路線、地圖圖層、標記歷史", "SCC/TL 維護；UCC 讀取彙整"],
        ["Tasking", "任務、指派、狀態、完成條件、照片/語音/位置證據", "SCC/TL 維護；TE/VO 回報"],
        ["Personnel / Resources", "人員狀態、裝備、電量、通訊狀態、可用性、調度紀錄", "SCC/Logistics 維護；UCC 彙整"],
        ["Medical", "傷患 ID、START/檢傷、生命徵象、後送、交接、醫院容量", "EMT 維護；SCC/UCC 只看必要營運摘要"],
        ["Comms / Alerts", "群組訊息、語音回報、警報、確認狀態、廣播範圍", "SCC/UCC 發布；各角色回覆確認"],
        ["Audit / AAR", "誰、何時、在哪裡、做了什麼、依據哪個任務或指令", "所有端寫入；UCC/SCC 匯出復盤"],
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    widths = [1900, 4750, 2710]
    set_table_widths(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            if r_idx == 0:
                run.bold = True
                set_cell_shading(cell, "F2F4F7")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 0:
                run.bold = True
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 16)
        style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("LinkGuard 真實救災架構查核報告")
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.bold = True
    r.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("依 INSARAG / FEMA ICS-NIMS / WHO EMT 官方資料整理｜2026-05-21")
    run.font.color.rgb = RGBColor(85, 85, 85)
    run.font.size = Pt(10)

    add_callout(
        doc,
        "結論：",
        "LinkGuard v0.3 的大方向正確，應保留「角色分流 + shared core + 離線同步 + AAR」；真正要補強的是實際現場權威分層、正式同步服務、裝置 provisioning、醫療資料邊界，以及 Mac UCC/SCC 從 v0.2 HQ 外殼逐步產品化。",
    )

    add_heading(doc, "1. 查核範圍與判斷基準")
    add_body(
        doc,
        "本文件不是要改程式碼，而是先把真實救災/USAR/ICS/EMT 場景中，系統架構應該遵守的角色、資料、同步與權限邊界整理清楚。判斷基準採官方一手資料：INSARAG Guidelines 2020、FEMA NIMS/ICS、FEMA Incident Action Planning、WHO EMT minimum standards。",
    )
    add_bullets(
        doc,
        [
            "INSARAG 觀點：大型結構倒塌災害需要 USAR 方法論、現場協調、UCC 與資訊管理。",
            "ICS/NIMS 觀點：現場管理應有共同階層，依需要展開 Command、Operations、Planning、Logistics、Finance/Admin。",
            "IAP 觀點：作戰不是散亂事件流，而是以 operational period、目標、資源指派與持續情資更新形成節奏。",
            "EMT 觀點：醫療隊伍需要與整體應變系統協作，但臨床資料與醫療流程不能被搜救任務流程淹沒。",
        ],
    )

    add_heading(doc, "2. 真實現場架構應該長什麼樣")
    add_body(
        doc,
        "真實架構不應是「所有資料都回到一台 Mac HQ」或「所有人看到所有功能」。較合理的模型是：現場戰術權威在 SCC/ICP，跨區與外部協調權威在 UCC/EOC/OSOCC 類角色，醫療資料由 EMT/Medical branch 維護，前線端只做低負擔回報與任務執行。",
    )
    add_bullets(
        doc,
        [
            "UCC 是跨區/戰略/對外協調層：看全局、管同步、做資源與資訊管理，不直接替每個 worksite 下細節戰術命令。",
            "SCC 是現場戰術層：負責分區、任務、現場安全、人員狀態與當下作戰圖。",
            "TL 是小隊/分區執行管理層：把 SCC 目標轉成可執行任務，維護 worksite 和回報狀態。",
            "TE/VO 是低負擔回報層：三秒操作、照片/位置/SOS/短訊息優先，不暴露複雜指揮台。",
            "EMT 是獨立醫療流程：檢傷、生命徵象、後送、交接與醫院容量，向 SCC/UCC 只分享營運必要摘要。",
        ],
    )

    add_heading(doc, "3. LinkGuard 角色對照表")
    add_matrix_table(doc)

    add_heading(doc, "4. 真實資料架構")
    add_body(
        doc,
        "資料架構應採「共同 incident state + 角色授權視圖 + domain authority」的設計。換句話說，所有版本使用同一套 ID、封包、時間戳、來源、權限與 audit contract；但不同角色只能維護自己負責的資料域。",
    )
    add_data_table(doc)

    add_heading(doc, "5. 同步與離線架構")
    add_body(
        doc,
        "真實現場最常遇到的是弱網、斷線、裝置電量低與多人同時回報。因此同步應是 offline-first，而不是依賴即時連線才能操作。每個端都應有本地 outbox、idempotency key、retry policy、明確 receipt，以及衝突處理策略。",
    )
    add_numbered(
        doc,
        [
            "端上先寫本地 snapshot 與 outbox，使用者操作立即完成。",
            "恢復連線後依 priority flush：SOS/安全/醫療 > 任務 > 照片/附件 > 一般聊天。",
            "SCC 作為現場 tactical authority，接收 field operations 與安全資料。",
            "UCC 作為 strategic/coordination authority，接收 SCC 彙整資料與跨區狀態。",
            "醫療 clinical payload 只在 EMT/醫療授權端完整流動；SCC/UCC 接收後送與容量摘要。",
            "所有 envelope 必須保留 source app、device ID、role、created time、received time、location、attachment references、ack state。",
        ],
    )

    add_heading(doc, "6. 模組切分建議")
    add_bullets(
        doc,
        [
            "Command：incident objectives、命令、警報、確認、跨中心廣播。",
            "Operations：sector/sub-sector/worksite、tasking、搜救進度、危險標記。",
            "Planning：IAP、operational period、情資彙整、資源預測、狀態板。",
            "Logistics/Comms：裝備、電力、通訊路由、LoRa/BLE gateway、補給與運輸。",
            "Medical：患者、START/檢傷、生命徵象、CCP、後送、醫院容量、交接紀錄。",
            "Finance/Admin：採購、工時、成本、文件、災後請款與行政軌跡。",
            "AAR/Audit：不可事後補寫的事件軌跡，支援時間線、決策紀錄、證據包與匯出。",
        ],
    )

    add_heading(doc, "7. 不建議的做法")
    add_bullets(
        doc,
        [
            "不要讓 Mac HQ 成為唯一真相來源；現場斷線時 SCC 與 field apps 仍要能獨立作業。",
            "不要讓所有角色看到所有資料；VO/TE 不應看到完整醫療病歷或指揮後台。",
            "不要把 AI 放在第一優先；現場優先是穩定、離線、低誤觸、三秒操作與可追蹤。",
            "不要只做 UI 分版，底層資料契約卻分裂；所有 app 應共享 ID、sync envelope、audit contract。",
            "不要把照片、語音、GPS、NFC、地圖當成 demo 欄位；這些都需要真實 device permission、背景限制與失敗狀態。",
            "不要忽略 device provisioning；每台裝置必須知道 incident、role、權限、同步 endpoint、撤權與遺失處理。",
        ],
    )

    add_heading(doc, "8. 建議落地順序")
    add_numbered(
        doc,
        [
            "先固定 branch/version policy：V0.2 維護線、V0.3 重建線、每次產品變更必須 bump version。",
            "建立正式 sync server endpoint 與 device provisioning，讓 FieldUI 不只本地 outbox。",
            "把 SCC 當現場 tactical hub 做完：incident、sector、worksite、task、safety、personnel status。",
            "把 EMT 醫療流程做成獨立閉環：patient card、START、vitals、evacuation、handover、hospital capacity。",
            "逐步替換 Mac UCC/SCC 的 v0.2 HQ 外殼，改成 v0.3 專用 strategic/tactical dashboard。",
            "接真實 GPS、相機、語音、NFC、地圖 SDK、通知、背景上傳與附件儲存。",
            "用 field drill 驗證：弱網、離線、多人同步、SOS、醫療後送、AAR 匯出。",
        ],
    )

    add_heading(doc, "9. 驗收清單")
    add_bullets(
        doc,
        [
            "每個 app 有清楚 role home，不只是同一套 UI 換名稱。",
            "每個 action 都能離線排隊，恢復後可追蹤 delivered/failed。",
            "SCC 和 UCC 的資料權威不同，但能同步同一個 incident state。",
            "EMT clinical data 有獨立權限邊界，只分享必要營運摘要。",
            "AAR 能重建時間線：人、裝置、角色、位置、任務、附件、決策與確認。",
            "GitHub 上每次產品 push/PR 都要求版本檔同步更新。",
        ],
    )

    add_heading(doc, "10. 來源")
    for idx, source in enumerate(SOURCES, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{idx}. {source['name']}").bold = True
        p.add_run(f" — {source['publisher']}. ")
        add_hyperlink(p, source["url"], source["url"])
        p2 = doc.add_paragraph(f"用途：{source['use']}")
        p2.paragraph_format.left_indent = Inches(0.25)
        p2.paragraph_format.space_after = Pt(4)

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("LinkGuard real architecture research | 2026-05-21")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
